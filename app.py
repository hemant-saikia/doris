import streamlit as st
import os
import re
from dotenv import load_dotenv
from urllib.parse import urlparse

load_dotenv()

PROVIDER_MAP = {
    "groq": {
        "label": "Groq",
        "env_key": "GROQ_API_KEY",
        "model": "llama-3.1-8b-instant",
    },
    "claude": {
        "label": "Claude (Anthropic)",
        "env_key": "CLAUDE_API_KEY",
        "model": "claude-sonnet-4-20250514",
    },
    "openrouter": {
        "label": "OpenRouter",
        "env_key": "OPENROUTER_API_KEY",
        "model": "openrouter/free",
    },
    "mistral": {
        "label": "Mistral",
        "env_key": "MISTRAL_API_KEY",
        "model": "mistral-large-latest",
    },
}


def fetch_url(url: str) -> str:
    import requests
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                      "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    }
    resp = requests.get(url, headers=headers, timeout=15)
    resp.raise_for_status()
    return _extract_text(resp.text)


def markdown_to_pdf(md: str, title: str) -> bytes:
    from fpdf import FPDF
    import io

    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=20)
    pw = pdf.epw

    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, title, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    def write(text: str, size: int = 11, bold: bool = False):
        pdf.set_font("Helvetica", "B" if bold else "", size)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(pw, size * 0.45, text)

    for line in md.split("\n"):
        stripped = line.strip()
        if not stripped:
            pdf.ln(3)
        elif stripped.startswith("# "):
            write(stripped[2:], 14, True)
            pdf.ln(2)
        elif stripped.startswith("## "):
            write(stripped[3:], 12, True)
            pdf.ln(2)
        elif stripped.startswith("### "):
            write(stripped[4:], 11, True)
            pdf.ln(1)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            write("  - " + stripped[2:])
        else:
            write(stripped)

    return io.BytesIO(pdf.output())


def _extract_text(html: str) -> str:
    text = re.sub(r"<script[^>]*>.*?</script>", "", html, flags=re.DOTALL | re.IGNORECASE)
    text = re.sub(r"<style[^>]*>.*?</style>", "", text, flags=re.DOTALL | re.IGNORECASE)
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"\n\s*\n", "\n", text)
    return text.strip()[:15000]


st.set_page_config(page_title="Job Application Assistant", page_icon="💼", initial_sidebar_state="expanded")

st.title("💼 Smart Job Application Assistant")

with st.sidebar:
    provider_name = st.selectbox(
        "LLM Provider",
        options=list(PROVIDER_MAP.keys()),
        format_func=lambda x: PROVIDER_MAP[x]["label"],
    )

    provider_info = PROVIDER_MAP[provider_name]
    model = provider_info["model"]

    st.caption(f"Model: `{model}`")

tab1, tab2, tab3 = st.tabs(["📝 Create Application", "📊 Track Applications", "⚙️ Settings"])

with tab1:
    st.header("Create New Application Package")

    col1, col2 = st.columns(2)

    with col1:
        st.subheader("Job Posting")

        input_mode = st.radio(
            "Input mode",
            ["Paste text", "URL"],
            horizontal=True,
            label_visibility="collapsed",
        )

        job_posting = ""

        if input_mode == "Paste text":
            job_posting = st.text_area(
                "Paste the job posting here",
                height=300,
                placeholder="Copy-paste the entire job description...",
            )
        else:
            job_posting = st.text_input(
                "Job posting URL",
                placeholder="https://...",
            )

    with col2:
        st.subheader("Your Base Resume")
        resume_file = st.file_uploader("Upload resume (PDF or DOCX)", type=["pdf", "docx"])

        base_resume = None
        if resume_file:
            if resume_file.type == "application/pdf":
                import PyPDF2
                pdf = PyPDF2.PdfReader(resume_file)
                base_resume = "\n".join([page.extract_text() for page in pdf.pages])
            else:
                from docx import Document
                doc = Document(resume_file)
                base_resume = "\n".join([p.text for p in doc.paragraphs])

            st.success(f"✅ Resume loaded ({len(base_resume)} characters)")

    can_generate = bool(job_posting and base_resume)

    if st.button("🚀 Generate Application Package", type="primary", disabled=not can_generate):
        try:
            api_key = os.getenv(provider_info["env_key"], "")
            if not api_key:
                st.error(f"❌ {provider_info['env_key']} not set in .env file")
                st.stop()

            status = st.status("Starting pipeline...", expanded=True)

            resolved_posting = job_posting
            if input_mode == "URL":
                status.write("🌐 **Fetching web page...**")
                if not urlparse(job_posting).scheme:
                    st.error("Invalid URL")
                    st.stop()
                resolved_posting = fetch_url(job_posting)

            status.write("🤖 **Extracting job data from contents...**")
            from llm import LLMConfig, create_llm
            config = LLMConfig(api_key=api_key, model=model)
            llm = create_llm(provider_name, config)
            from agents.job_analyzer import JobAnalyzerAgent
            job_analyzer = JobAnalyzerAgent(llm)
            analysis = job_analyzer.analyze(resolved_posting)

            status.write("✍️ **Writing new resume with AI...**")
            from agents.resume_optimizer import ResumeOptimizerAgent
            resume_optimizer = ResumeOptimizerAgent(llm)
            optimized_resume = resume_optimizer.optimize(base_resume, analysis)

            status.write("💌 **Writing cover letter with AI...**")
            from agents.cover_letter import CoverLetterAgent
            cover_letter_agent = CoverLetterAgent(llm)
            cover_letter = cover_letter_agent.generate(analysis, optimized_resume)

            status.update(state="complete", label="✅ Application package ready!")

            st.subheader("📋 Job Analysis")
            st.json({
                "role": analysis.role,
                "company": analysis.company,
                "required_skills": analysis.required_skills,
                "ats_keywords": analysis.ats_keywords,
            })

            st.subheader("📄 Optimized Resume")
            st.markdown(optimized_resume)
            st.download_button(
                "⬇️ Download Resume (PDF)",
                markdown_to_pdf(optimized_resume, f"Resume - {analysis.role}"),
                f"resume_{analysis.company}.pdf",
                mime="application/pdf",
            )

            st.subheader("✉️ Cover Letter")
            st.markdown(cover_letter)
            st.download_button(
                "⬇️ Download Cover Letter (PDF)",
                markdown_to_pdf(cover_letter, f"Cover Letter - {analysis.role}"),
                f"cover_letter_{analysis.company}.pdf",
                mime="application/pdf",
            )

        except Exception as e:
            st.error(f"❌ Error: {str(e)}")
            st.info("Check your API key in .env and try again")

with tab2:
    st.header("📊 Application Tracker")
    st.info("Coming soon - track all your applications in one place")

with tab3:
    st.header("⚙️ Settings")
    tone = st.selectbox("Cover letter tone", ["professional", "enthusiastic", "confident"])
    st.info("Settings will be saved in future versions")
